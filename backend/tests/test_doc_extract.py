import io
import os
import zipfile

import pytest

for _k, _v in {"MONGO_URL": "mongodb://unused", "DB_NAME": "t", "MEDIA_DIR": "/tmp/pp-test-media", "JWT_SECRET": "x" * 40}.items():
    os.environ.setdefault(_k, _v)

from docx import Document  # noqa: E402

import doc_extract as de  # noqa: E402


def test_markdown_and_text():
    md = "# Planning Template\r\n\r\n\r\n\r\n- Hook\n- Scripture\n"
    out = de.extract("Template.MD", md.encode())
    assert out.startswith("# Planning Template") and "\n\n\n" not in out and "\r" not in out
    assert de.extract("a.txt", "café ☕".encode("utf-8")) == "café ☕"
    assert de.extract("a.txt", "café".encode("cp1252")) == "café"           # Windows Notepad files
    assert de.extract("a.txt", "héllo".encode("utf-16")) == "héllo"


def test_docx_still_works():
    d = Document()
    d.add_heading("Title", 1)
    d.add_paragraph("Body text")
    buf = io.BytesIO()
    d.save(buf)
    out = de.extract("x.docx", buf.getvalue())
    assert "# Title" in out and "Body text" in out


def make_pdf(text: str) -> bytes:
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"]
    stream = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET"
    objs += [f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxml\n{xref}\n%%EOF".replace("startxml", "startxref").encode()


def test_pdf():
    assert "Hello planning" in de.extract("t.pdf", make_pdf("Hello planning"))


def test_rtf_html_odt():
    assert "Plain words" in de.extract("t.rtf", rb"{\rtf1\ansi {\b Plain words} here\par next}")
    html = b"<html><head><style>x{}</style></head><body><h1>Head</h1><p>One</p><p>Two<br>three</p><script>bad()</script></body></html>"
    out = de.extract("t.html", html)
    assert "Head" in out and "One" in out and "bad()" not in out and "x{}" not in out
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("content.xml", '<o:document xmlns:o="urn:o" xmlns:t="urn:t"><t:p>First</t:p><t:h>Second</t:h></o:document>')
    assert de.extract("t.odt", buf.getvalue()) == "First\n\nSecond"


def test_rejects_bad_input():
    with pytest.raises(de.UnsupportedDocument):
        de.extract("a.exe", b"MZ")
    with pytest.raises(de.UnsupportedDocument):
        de.extract("a.txt", b"abc\x00\x01\x02" * 100)
    with pytest.raises(Exception):
        de.extract("a.pdf", b"not a pdf")
    assert de.extract("empty.pdf", make_pdf("")) == ""
