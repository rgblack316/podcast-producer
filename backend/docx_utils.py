import io
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches


def extract_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            level = re.sub(r"\D", "", style) or "2"
            lines.append(f"{'#' * min(int(level), 4)} {t}")
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n\n".join(lines)


def _add_runs(par, text: str):
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            run = par.add_run(chunk.replace("*", "") if i % 2 == 0 else chunk)
            run.bold = i % 2 == 1


def markdown_to_docx(title: str, meta: list, md: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)
    for line in meta:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.italic = True
        r.font.color.rgb = RGBColor(0x57, 0x53, 0x4E)
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 4))
        elif line.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(line.lstrip()[1:].strip().replace("**", ""))
            r.italic = True
            r.font.color.rgb = RGBColor(0x78, 0x35, 0x0F)
        elif re.match(r"^\s*[-*+]\s+", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        elif re.match(r"^\s*(---|\*\*\*)\s*$", line):
            doc.add_paragraph("")
        else:
            _add_runs(doc.add_paragraph(), line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
